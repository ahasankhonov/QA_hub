#!/usr/bin/env python3
"""R-11 — test-case-list report derived from the run ledger + analyst plans. No hand-authored verdicts.
usage: python3 scripts/gen-testcases.py --project <slug> --run <runId> [--out <file.docx>]
Writes QA-TestCases-<run>.docx (landscape table, one row per case) and _TESTCASES-<run>.md (grep-able twin)."""
import json, sys, re, glob, os, datetime
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

a = dict(zip(sys.argv[1::2], sys.argv[2::2]))
slug, run = a.get('--project'), a.get('--run')
if not slug or not run: sys.exit(__doc__)
HUB = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
rd = os.path.join(HUB, 'projects', slug, 'qa-reports')
ledger = json.load(open(os.path.join(rd, f'_LEDGER-{run}.json')))
plans = {}
for f in glob.glob(os.path.join(rd, f'_PLAN-{run}-*.json')):
    p = json.load(open(f))
    for s in p.get('steps', []): plans[s['caseId']] = s
out = os.path.join(HUB, a['--out']) if a.get('--out') else os.path.join(rd, f'QA-TestCases-{run}.docx')
md_out = os.path.join(rd, f'_TESTCASES-{run}.md')

def plan_for(cid):
    if cid in plans: return plans[cid]
    # per-plan expansions: 708-P2-6-merchant-submit → 708-P<n>-6 (analyst wrote one template step per plan slot)
    g = re.sub(r'-P\d+-', '-P<n>-', re.sub(r'-[a-z][a-z-]*$', '', cid))
    if g in plans: return plans[g]
    # executor-expanded ids (e.g. 708-P1-3 → 708-3, 704-D2 → 704-2): fall back to the base step
    m = re.match(r'^(\d+)-(?:[A-Z]+\d*-)?([A-Z]?)(\d+)$', cid)
    if m:
        for cand in (f'{m.group(1)}-{m.group(3)}', f'{m.group(1)}-{m.group(2)}{m.group(3)}'):
            if cand in plans: return plans[cand]
    return {}

STATUS_FILL = {'PASS': 'C6EFCE', 'FAIL': 'FFC7CE', 'NOT_TESTED': 'E7E6E6', 'NEEDS_MANUAL': 'FFEB9C', 'BLOCKED': 'FFEB9C'}
def shade(cell, hex_):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), hex_); tcPr.append(shd)
def put(cell, text, bold=False, size=7.5):
    cell.text = ''
    p = cell.paragraphs[0]; r = p.add_run(str(text)); r.font.size = Pt(size); r.bold = bold
    p.paragraph_format.space_after = Pt(0)
def steps_text(s):
    acts = s.get('actions') or []
    return '\n'.join(f'{i+1}. {x}' for i, x in enumerate(acts)) if acts else '—'
def actual_text(c):
    note = (c.get('note') or '').strip()
    ev = c.get('evidence') or {}
    extra = []
    if isinstance(ev, dict):
        if ev.get('httpStatus'): extra.append(f"HTTP {ev['httpStatus']}")
        if ev.get('screenshot'): extra.append(f"screenshot: {os.path.basename(str(ev['screenshot']))}")
    return (note or '—') + (('\n' + ' · '.join(extra)) if extra else '')

doc = Document()
sec = doc.sections[0]; sec.orientation = WD_ORIENT.LANDSCAPE
sec.page_width, sec.page_height = sec.page_height, sec.page_width
for side in ('left_margin', 'right_margin', 'top_margin', 'bottom_margin'): setattr(sec, side, Inches(0.5))
st = doc.styles['Normal']; st.font.name = 'Calibri'; st.font.size = Pt(8)

tot = {}
for t in ledger['tickets']:
    for c in t['cases']: tot[c['verdict']] = tot.get(c['verdict'], 0) + 1
ncases = sum(tot.values())
doc.add_heading(f'{slug.capitalize()} QA — Test Case List — run {run}', 0)
p = doc.add_paragraph(); p.add_run(f"Build: {ledger.get('build') or ledger['tickets'][0].get('buildId')}\n").bold = True
p.add_run(f"Generated {datetime.date.today().isoformat()} from _LEDGER-{run}.json + _PLAN-{run}-*.json (R-11: verdicts and counts derived, not typed).\n")
p.add_run(f"{len(ledger['tickets'])} tickets · {ncases} cases · " + ' · '.join(f'{k} {v}' for k, v in sorted(tot.items())))

# summary table
sm = doc.add_table(rows=1, cols=7); sm.style = 'Light Grid Accent 1'
for i, h in enumerate(['Ticket', 'Verdict', 'Cases', 'PASS', 'FAIL', 'NOT_TESTED', 'NEEDS_MANUAL']): put(sm.rows[0].cells[i], h, True, 8)
for t in ledger['tickets']:
    vc = {}
    for c in t['cases']: vc[c['verdict']] = vc.get(c['verdict'], 0) + 1
    row = sm.add_row().cells
    verdict = t['verdict'] + (f":{t['partialReason']}" if t.get('partialReason') else '')
    for i, v in enumerate([t['id'], verdict, len(t['cases']), vc.get('PASS', 0), vc.get('FAIL', 0), vc.get('NOT_TESTED', 0), vc.get('NEEDS_MANUAL', 0)]): put(row[i], v, size=8)

md = [f'# {slug} {run} — test case list', '', f"Build: {ledger.get('build')}", '']
HEAD = ['#', 'Case ID', 'Test case', 'Surface / actor', 'Steps', 'Expected', 'Actual (recorded)', 'Status']
W = [0.3, 0.7, 1.4, 1.1, 2.6, 1.7, 2.1, 0.8]
n = 0
for t in ledger['tickets']:
    plan_title = next((p for p in plans.values() if p['caseId'].startswith(t['id'].split('-')[1] + '-')), None)
    doc.add_heading(f"{t['id']} — verdict {t['verdict']}" + (f" ({t['partialReason']})" if t.get('partialReason') else ''), 1)
    if t.get('comment'):
        q = doc.add_paragraph(str(t['comment'])[:600]); q.runs[0].font.size = Pt(8)
    md += [f"## {t['id']} — {t['verdict']}", '', '| # | Case | Title | Steps | Expected | Actual | Status |', '|---|---|---|---|---|---|---|']
    tb = doc.add_table(rows=1, cols=len(HEAD)); tb.style = 'Table Grid'
    for i, h in enumerate(HEAD): put(tb.rows[0].cells[i], h, True); shade(tb.rows[0].cells[i], 'D9E2F3')
    for c in t['cases']:
        n += 1; s = plan_for(c['caseId'])
        vals = [n, c['caseId'], s.get('title') or '—', f"{c.get('surface') or s.get('surface') or '—'}\n{s.get('actor') or ''}".strip(),
                steps_text(s), s.get('expected') or '—', actual_text(c), c['verdict']]
        row = tb.add_row().cells
        for i, v in enumerate(vals): put(row[i], v, bold=(i == 7))
        shade(row[7], STATUS_FILL.get(c['verdict'], 'FFFFFF'))
        md.append('| ' + ' | '.join(str(v).replace('\n', '<br>').replace('|', '\\|') for v in [n, c['caseId'], vals[2], vals[4], vals[5], (c.get('note') or '—'), c['verdict']]) + ' |')
    for row in tb.rows:
        for i, w in enumerate(W): row.cells[i].width = Inches(w)
    md.append('')
doc.save(out); open(md_out, 'w').write('\n'.join(md))
print(f'[gen-testcases] wrote {out} and {md_out} — {len(ledger["tickets"])} tickets, {n} cases ({", ".join(f"{k} {v}" for k, v in sorted(tot.items()))}); plan match {sum(1 for t in ledger["tickets"] for c in t["cases"] if plan_for(c["caseId"]))}/{n}')
